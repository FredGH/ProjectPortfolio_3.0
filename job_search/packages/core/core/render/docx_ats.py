"""The locked ATS .docx writer. It only ever adds plain paragraphs, headings
and list bullets, so tables, text boxes, headers/footers and images have no
code path: the ATS rules are enforced by construction, not by prompting."""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.document import Document as DocumentType
from docx.shared import Mm, Pt, RGBColor

from core.render.model import Block, RenderDoc

_FONT = "Calibri"
_BODY_PT = 10.5


def _style_document(document: DocumentType) -> None:
    """Set A4 single-column page, margins and base fonts.

    Args:
        document: The new document to configure.
    """
    section = document.sections[0]
    section.page_width = Mm(210)
    section.page_height = Mm(297)
    for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(section, side, Mm(18))
    normal = document.styles["Normal"]
    normal.font.name = _FONT
    normal.font.size = Pt(_BODY_PT)
    normal.paragraph_format.space_after = Pt(2)
    heading = document.styles["Heading 1"]
    heading.font.name = _FONT
    heading.font.size = Pt(12)
    heading.font.bold = True
    heading.font.color.rgb = RGBColor(0, 0, 0)
    heading.paragraph_format.space_before = Pt(10)
    heading.paragraph_format.space_after = Pt(3)


def _add_block(document: DocumentType, block: Block) -> None:
    """Append one block as a plain paragraph.

    Args:
        document: The document being built.
        block: The block to write.
    """
    if block.kind == "heading":
        document.add_paragraph(block.text, style="Heading 1")
        return
    if block.kind == "bullet":
        document.add_paragraph(block.text, style="List Bullet")
        return
    paragraph = document.add_paragraph()
    run = paragraph.add_run(block.text)
    if block.kind == "name":
        run.bold = True
        run.font.size = Pt(20)
    elif block.kind == "headline":
        run.bold = True
        run.font.size = Pt(12)
    elif block.kind == "role_title":
        run.bold = True
        paragraph.paragraph_format.space_before = Pt(6)
        paragraph.paragraph_format.keep_with_next = True
    elif block.kind == "role_meta":
        run.italic = True
        paragraph.paragraph_format.keep_with_next = True


def write_docx(doc: RenderDoc, path: Path) -> None:
    """Write the ATS-safe .docx.

    Args:
        doc: The render model.
        path: Where to save the file.
    """
    document = Document()
    _style_document(document)
    for block in doc.blocks:
        _add_block(document, block)
    document.core_properties.title = doc.title[:255]
    document.core_properties.author = ""
    document.save(str(path))
