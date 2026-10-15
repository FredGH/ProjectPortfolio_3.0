"""Unit tests for rendering letters and pitches to files."""

from __future__ import annotations

import tempfile
import unittest
from pathlib import Path

from docx import Document
from pypdf import PdfReader
from tests.artefacts_fixtures import (
    TODAY,
    letter_reply_json,
    make_pack,
    pitch_reply_json,
)

from core.artefacts.assemble import build_letter, build_pitch, parse_writer_reply
from core.artefacts.render import (
    RenderError,
    letter_download_bundle,
    letter_render_doc,
    pitch_download_bundle,
    render_cover_letter_files,
    render_pitch_file,
)
from core.render.pdf_designed import HELVETICA


def _letter():  # type: ignore[no-untyped-def]
    """Assemble the fixture letter.

    Returns:
        A `CoverLetter`.
    """
    return build_letter(
        make_pack(),
        parse_writer_reply(letter_reply_json(), "cover_letter"),
        today=TODAY,
    )


class TestLetterFiles(unittest.TestCase):
    """The letter is written as docx, txt and pdf and verified."""

    def setUp(self) -> None:
        """Create a scratch directory."""
        self._tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self._tmp.cleanup)
        self.dir = Path(self._tmp.name)

    def test_the_files_are_named_for_the_surname_title_and_company(self) -> None:
        files = render_cover_letter_files(_letter(), self.dir, fonts=HELVETICA)
        self.assertEqual(
            files.docx_path.name,
            "Fixture_Data_Engineer_Zzfixture_Intel_Ltd_cover_letter.docx",
        )
        self.assertEqual(files.txt_path.suffix, ".txt")
        assert files.pdf_path is not None
        self.assertEqual(files.pdf_path.suffix, ".pdf")

    def test_the_docx_text_matches_the_txt_twin_and_has_the_exact_title_twice(
        self,
    ) -> None:
        files = render_cover_letter_files(_letter(), self.dir, fonts=HELVETICA)
        docx_text = "\n".join(p.text for p in Document(str(files.docx_path)).paragraphs)
        self.assertEqual(
            docx_text.strip().split(), files.txt_path.read_text().strip().split()
        )
        self.assertGreaterEqual(docx_text.count("Data Engineer"), 2)
        self.assertIn("Re: Data Engineer", docx_text)

    def test_the_pdf_reads_back_with_the_exact_title(self) -> None:
        files = render_cover_letter_files(_letter(), self.dir, fonts=HELVETICA)
        assert files.pdf_path is not None
        text = "".join(PdfReader(str(files.pdf_path)).pages[0].extract_text().split())
        self.assertIn("Re:DataEngineer", text)

    def test_ats_only_skips_the_pdf(self) -> None:
        files = render_cover_letter_files(_letter(), self.dir, with_pdf=False)
        self.assertIsNone(files.pdf_path)
        self.assertEqual(
            sorted(p.suffix for p in self.dir.iterdir()), [".docx", ".txt"]
        )

    def test_a_letter_that_loses_its_title_fails_and_leaves_no_files(self) -> None:
        letter = _letter()
        letter.reference = "Re: Something Else"
        letter.paragraphs[0].sentences[0].text = "I am applying."
        with self.assertRaises(RenderError):
            render_cover_letter_files(letter, self.dir, fonts=HELVETICA)
        self.assertEqual(list(self.dir.iterdir()), [])

    def test_the_render_doc_is_one_plain_column(self) -> None:
        doc = letter_render_doc(_letter())
        kinds = {b.kind for b in doc.blocks}
        self.assertLessEqual(kinds, {"name", "contact", "headline", "paragraph"})
        self.assertEqual(doc.blocks[0].text, "Zz Fixture")
        self.assertEqual(doc.title, "Data Engineer")


class TestPitchFile(unittest.TestCase):
    """The pitch is a plain-text script of the three variants."""

    def test_the_script_has_three_headed_variants_each_starting_with_the_title(
        self,
    ) -> None:
        pitch = build_pitch(
            make_pack(), parse_writer_reply(pitch_reply_json(), "pitch")
        )
        with tempfile.TemporaryDirectory() as tmp:
            path = render_pitch_file(pitch, Path(tmp))
            text = path.read_text()
        self.assertEqual(
            path.name, "Fixture_Data_Engineer_Zzfixture_Intel_Ltd_pitch.txt"
        )
        for seconds in (30, 60, 90):
            self.assertIn(f"{seconds} seconds", text)
        self.assertEqual(text.count("the Data Engineer role"), 3)

    def test_a_pitch_variant_without_the_title_fails(self) -> None:
        pitch = build_pitch(
            make_pack(), parse_writer_reply(pitch_reply_json(), "pitch")
        )
        pitch.variants[1].sentences[0].text = "Hello."
        with tempfile.TemporaryDirectory() as tmp, self.assertRaises(RenderError):
            render_pitch_file(pitch, Path(tmp))


class TestDownloadBundles(unittest.TestCase):
    """The UI gets the files as bytes, rendered in memory."""

    def test_the_letter_bundle_has_pdf_docx_and_txt(self) -> None:
        bundle = letter_download_bundle(_letter(), fonts=HELVETICA)
        names = [name for name, _mime, _data in bundle]
        self.assertEqual([n.rsplit(".", 1)[1] for n in names], ["pdf", "docx", "txt"])
        self.assertTrue(all(data for _n, _m, data in bundle))

    def test_the_pitch_bundle_is_the_text_script(self) -> None:
        pitch = build_pitch(
            make_pack(), parse_writer_reply(pitch_reply_json(), "pitch")
        )
        ((name, mime, data),) = pitch_download_bundle(pitch)
        self.assertTrue(name.endswith("_pitch.txt"))
        self.assertEqual(mime, "text/plain")
        self.assertIn(b"30 seconds", data)


if __name__ == "__main__":
    unittest.main()
